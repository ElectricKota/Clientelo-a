"""Exporty: CSV/XLSX vzorku, všech kandidátů, dokladů a ověřovací formulář (PLAN.md §9)."""

from __future__ import annotations

import csv
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import TIER_A, Evidence, Investor

SEED = 20261011
TIER_ORDER = {"A": 0, "B": 1, "C": 2}


def _eur(v: float | None) -> str:
    return "" if v is None else f"{v:,.0f}".replace(",", " ")


def primary_evidence(i: Investor, cutoff: str) -> Evidence | None:
    inv = i.investments(cutoff) or i.investments()
    if not inv:
        prof = [e for e in i.evidence if e.claim == "profile"]
        return prof[0] if prof else None
    return sorted(inv, key=lambda e: (TIER_ORDER[e.tier],
                                      -int(e.published_date.replace("-", "") or 0)))[0]


def investor_row(i: Investor, cutoff: str) -> dict:
    inv = i.investments(cutoff)
    pe = primary_evidence(i, cutoff)
    return {
        "investor": i.name,
        "typ": i.type,
        "tagy": ",".join(sorted(i.tags)),
        "status": i.status,
        "duvod_kod": i.exclusion_code,
        "duvody": " | ".join(i.reasons),
        "spravce": i.manager_name,
        "spravce_ico": i.manager_ico,
        "spravce_zdroj": i.manager_source,
        "sektory": ",".join(i.sectors),
        "faze": ",".join(i.stages),
        "ticket_min_eur": _eur(i.ticket_min_eur),
        "ticket_max_eur": _eur(i.ticket_max_eur),
        "ticket_zdroj": i.ticket_source,
        "ticket_k_datu": i.ticket_as_of,
        "aum_eur": _eur(i.aum_eur),
        "aum_zdroj": i.aum_source,
        "aum_k_datu": i.aum_as_of,
        "web": i.website,
        "investic_v_okne": len(inv),
        "investic_tier_A": sum(1 for e in inv if e.tier == TIER_A),
        "posledni_investice": max((e.published_date for e in inv), default=""),
        "doklad_url": pe.source_url if pe else "",
        "doklad_datum": pe.published_date if pe else "",
        "doklad_tier": pe.tier if pe else "",
        "doklad_firma": pe.company if pe else "",
        "doklad_citace": pe.snippet if pe else "",
        "confidence": i.confidence,
        "zdroje_objeveni": ",".join(sorted(i.sources)),
    }


def _write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def _sheet(wb: Workbook, title: str, rows: list[dict], first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    if not rows:
        return
    ws.append(list(rows[0]))
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r.values()))
    for n, col in enumerate(rows[0], start=1):
        width = min(60, max(10, len(col) + 2, *(len(str(r[col])) for r in rows[:200])))
        ws.column_dimensions[get_column_letter(n)].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def evidence_rows(investors: list[Investor]) -> list[dict]:
    rows = []
    for i in sorted(investors, key=lambda x: x.name.lower()):
        if i.status not in ("included", "review"):
            continue
        for e in sorted(i.evidence, key=lambda e: (e.claim, e.published_date), reverse=True):
            rows.append({"investor": i.name, "status": i.status, "tvrzeni": e.claim, "hodnota": e.value,
                         "tier": e.tier, "typ_zdroje": e.source_type, "datum_dokladu": e.published_date,
                         "stazeno": e.retrieved_at, "url": e.source_url, "citace": e.snippet,
                         "firma": e.company, "castka_eur": _eur(e.amount_eur), "sha256": e.sha256})
    return rows


def export_all(investors: list[Investor], cutoff: str, out_dir: Path) -> dict[str, int]:
    sample = [i for i in investors if i.status == "included" and i.type in ("vc", "cvc")]
    sample.sort(key=lambda i: (-i.confidence, i.name.lower()))
    candidates = sorted(investors, key=lambda i: ({"included": 0, "review": 1, "excluded": 2}[i.status], i.name.lower()))
    rows_sample = [investor_row(i, cutoff) for i in sample]
    rows_all = [investor_row(i, cutoff) for i in candidates]
    rows_ev = evidence_rows(investors)
    _write_csv(out_dir / "investors_cz_vc.csv", rows_sample)
    _write_csv(out_dir / "all_candidates.csv", rows_all)
    _write_csv(out_dir / "evidence.csv", rows_ev)
    wb = Workbook()
    _sheet(wb, "VC vzorek CZ", rows_sample, first=True)
    _sheet(wb, "Všichni kandidáti", rows_all)
    _sheet(wb, "Doklady", rows_ev)
    wb.save(out_dir / "investors_cz_vc.xlsx")
    return {"sample": len(rows_sample), "candidates": len(rows_all), "evidence": len(rows_ev)}


def verification_form(investors: list[Investor], cutoff: str, path: Path,
                      n_included: int = 35, n_review: int = 5, n_excluded: int = 5) -> int:
    """Formulář pro ruční ověření (≤ ~40 min). Pořadí řádků je náhodné se zafixovaným seedem."""
    rnd = random.Random(SEED)
    inc = [i for i in investors if i.status == "included" and i.type in ("vc", "cvc")]
    if len(inc) > n_included:
        inc = rnd.sample(inc, n_included)
    rev = [i for i in investors if i.status == "review" or (i.status == "excluded" and i.investments())]
    rev = rnd.sample(rev, min(n_review, len(rev)))
    noev = [i for i in investors if i.exclusion_code == "NO_PUBLIC_EVIDENCE"
            and any(e.list_code == "ZISIF_PAR15" for e in i.jerrs)]
    noev = rnd.sample(noev, min(n_excluded, len(noev)))
    picked = [("included", i) for i in inc] + [("not_included", i) for i in rev + noev]
    rnd.shuffle(picked)
    rows = []
    for n, (group, i) in enumerate(picked, start=1):
        pe = primary_evidence(i, cutoff)
        claim = (f"typ={i.type}; sektory={','.join(i.sectors) or '-'}; "
                 f"ticket={_eur(i.ticket_min_eur)}–{_eur(i.ticket_max_eur)} EUR ({i.ticket_source or '-'}); "
                 f"AUM={_eur(i.aum_eur) or '-'} EUR ({i.aum_source or '-'} {i.aum_as_of})")
        rows.append({
            "row": n,
            "skupina": group,
            "investor": i.name,
            "pipeline_rika": ("ZAŘAZEN: " if group == "included" else f"NEZAŘAZEN ({i.status}/{i.exclusion_code}): ")
                             + (claim if group == "included" else " | ".join(i.reasons)),
            "cesky_spravce": f"{i.manager_name} (IČO {i.manager_ico})" if i.manager_ico else "",
            "ares_url": f"https://ares.gov.cz/ekonomicke-subjekty?ico={i.manager_ico}" if i.manager_ico else "",
            "doklad_url": pe.source_url if pe else "",
            "doklad_datum": pe.published_date if pe else "",
            "doklad_citace": (pe.snippet[:300] if pe else ""),
            "is_real_vc": "",
            "evidence_ok": "",
            "fields_ok": "",
            "note": "",
        })
    _write_csv(path, rows)
    return len(rows)
